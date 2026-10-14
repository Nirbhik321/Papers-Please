"""
exporter.py — Generate exports: PDF cheat sheet and the question bank (CSV and Excel).

PDF: an A4 cheat sheet — the most repeated topics per module and the marks they cover.
CSV / Excel: every printed question, grouped by topic, for teachers to build internal assessments.
"""

import csv
import io
import re
from datetime import date

from fpdf import FPDF, XPos, YPos

from modules.scorer import MAX_MODULE_MARKS


# ── ASCII sanitizer (Helvetica supports Latin-1 only) ─────────────────────────

_UNICODE_REPLACEMENTS = str.maketrans({
    "\u2018": "'",  "\u2019": "'",   # curly single quotes
    "\u201c": '"',  "\u201d": '"',   # curly double quotes  ← the current crash
    "\u2013": "-",  "\u2014": "-",   # en-dash, em-dash
    "\u2022": "*",  "\u00b7": "*",   # bullet points
    "\u00ae": "(R)", "\u00a9": "(C)", # registered, copyright
    "\u2026": "...",                  # ellipsis
    "\u00b0": "deg",                  # degree
    "\u2192": "->", "\u2190": "<-",  # arrows
    "\u2605": "*",  "\u2606": "*",   # filled/open stars
    "\u00e2": "a",  "\u00e9": "e",   # common accented chars from bad OCR
})


def _safe(text: str) -> str:
    """Strip non-Latin-1 characters so fpdf Helvetica never crashes."""
    text = text.translate(_UNICODE_REPLACEMENTS)
    # Replace any remaining non-Latin-1 chars with '?'
    return text.encode("latin-1", errors="replace").decode("latin-1")


# ── Palette (matches the website) ──────────────────────────────────────────────
INK        = (27,  34,  51)
INK_2      = (74,  81,  99)
LINE       = (221, 215, 201)
BAND       = (242, 239, 231)
STAMP      = (169, 50,  38)
AMBER      = (208, 138, 60)
OUTLINE    = (138, 142, 153)
OK_GREEN   = (30,  90,  57)
WHITE      = (255, 255, 255)

PAGE_W, MARGIN = 210, 15
CONTENT_W = PAGE_W - 2 * MARGIN
ROWS_PER_MODULE = 3


def generate_cheat_sheet(
    subject_name: str,
    subject_code: str,
    module_ladders: dict[int, list[dict]],   # {module_no: [step, ...]}
    total_papers: int,
    output_path: str,
    site_url: str = "",
) -> str:
    """
    Generate the one-page A4 cheat sheet: for each module, the topics that
    repeat most (up to the point where the module's marks are covered).

    Args:
        subject_name: e.g. "Computer Networks"
        subject_code: e.g. "BCS502"
        module_ladders: {module_no: result of scorer.build_marks_ladder()}
        total_papers: number of papers analysed
        output_path: file path to write the PDF
        site_url: optional website base, printed in the footer

    Returns:
        output_path (for convenience)
    """
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_margins(MARGIN, 14, MARGIN)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.set_title(_safe(f"{subject_code} {subject_name} - cheat sheet"))
    pdf.add_page()

    _draw_header(pdf, subject_name, subject_code, total_papers)
    for module_no in sorted(module_ladders):
        steps = module_ladders[module_no]
        if steps:
            _draw_module(pdf, module_no, steps, total_papers)
    _draw_footer(pdf, subject_code, site_url)

    pdf.output(output_path)
    return output_path


def _tier(step: dict, total_papers: int) -> tuple[tuple, tuple]:
    """(fill, border) of the frequency square: 5-6 of 6 papers, 3-4, 1-2."""
    pct = step["frequency_pct"]
    if total_papers > 1 and pct >= 0.8:
        return STAMP, STAMP
    if total_papers > 1 and pct >= 0.5:
        return AMBER, AMBER
    return WHITE, OUTLINE


def _square(pdf: FPDF, x: float, y: float, fill: tuple, border: tuple, size: float = 2.6):
    pdf.set_fill_color(*fill)
    pdf.set_draw_color(*border)
    pdf.set_line_width(0.35)
    pdf.rect(x, y, size, size, style="DF")


def _draw_header(pdf: FPDF, subject_name: str, subject_code: str, total_papers: int):
    top = pdf.get_y()
    pdf.set_text_color(*STAMP)
    pdf.set_font("Helvetica", "B", 7)
    pdf.cell(0, 4, "PAPERS PLEASE  -  CHEAT SHEET", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_text_color(*INK)
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(125, 9, _safe(f"{subject_code}  {subject_name}"[:60]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_text_color(*INK_2)
    pdf.set_font("Helvetica", "", 8)
    papers = f"{total_papers} paper{'s' if total_papers != 1 else ''}"
    pdf.cell(125, 5, f"From {papers}  -  generated {date.today().strftime('%d %b %Y')}",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    bottom = pdf.get_y()

    # Legend, top right
    legend = [(STAMP, STAMP, "Very likely  -  most papers"), (AMBER, AMBER, "Likely  -  half or more"),
              (WHITE, OUTLINE, "Occasional")]
    pdf.set_font("Helvetica", "", 7.5)
    pdf.set_text_color(*INK_2)
    for i, (fill, border, label) in enumerate(legend):
        y = top + 3 + i * 5
        _square(pdf, 150, y, fill, border)
        pdf.set_xy(154, y - 0.9)
        pdf.cell(40, 4.4, label)

    pdf.set_y(max(bottom, top + 18) + 1.5)
    pdf.set_draw_color(*INK)
    pdf.set_line_width(0.6)
    pdf.line(MARGIN, pdf.get_y(), PAGE_W - MARGIN, pdf.get_y())
    pdf.ln(3.5)


def _draw_module(pdf: FPDF, module_no: int, steps: list[dict], total_papers: int):
    full_at = next((s["rank"] for s in steps if s["full_coverage"]), None)
    shown = steps[: min(full_at or ROWS_PER_MODULE, ROWS_PER_MODULE)]
    last = shown[-1]
    if full_at and full_at <= ROWS_PER_MODULE:
        ladder = f"Top {full_at} -> full {MAX_MODULE_MARKS}M"
    else:
        ladder = f"Top {len(shown)} -> ~{last['cumulative_expected']:.0f} of {MAX_MODULE_MARKS}M"

    if pdf.get_y() > 297 - 16 - 30:   # keep a module's band with its first rows
        pdf.add_page()

    y = pdf.get_y()
    pdf.set_fill_color(*BAND)
    pdf.rect(MARGIN, y, CONTENT_W, 6.5, style="F")
    pdf.set_xy(MARGIN + 2.5, y + 1)
    pdf.set_text_color(*INK)
    pdf.set_font("Helvetica", "B", 9.5)
    pdf.cell(100, 4.5, f"Module {module_no}")
    pdf.set_text_color(*OK_GREEN)
    pdf.set_font("Helvetica", "B", 8)
    pdf.set_xy(MARGIN, y + 1)
    pdf.cell(CONTENT_W - 2.5, 4.5, ladder, align="R")
    pdf.set_y(y + 8)

    for step in shown:
        _draw_row(pdf, step, total_papers)
    pdf.ln(3)


def _draw_row(pdf: FPDF, step: dict, total_papers: int):
    if pdf.get_y() > 297 - 16 - 14:   # never split a row across pages
        pdf.add_page()
    y = pdf.get_y()
    fill, border = _tier(step, total_papers)
    _square(pdf, MARGIN + 2.5, y + 1.1, fill, border)

    text_x, text_w = MARGIN + 8, CONTENT_W - 32
    pdf.set_xy(text_x, y)
    pdf.set_text_color(*INK)
    pdf.set_font("Helvetica", "B", 9.5)
    pdf.cell(text_w, 4.8, _safe(step["topic_label"] or step["representative_text"][:60])[:80])

    # Right-hand numbers: papers and usual marks
    pdf.set_font("Helvetica", "", 8.5)
    pdf.set_text_color(*INK_2)
    pdf.set_xy(PAGE_W - MARGIN - 22, y)
    pdf.cell(11, 4.8, f"{step['frequency']}/{total_papers}", align="R")
    pdf.set_text_color(*INK)
    pdf.set_font("Helvetica", "B", 8.5)
    marks = int(round(step.get("avg_marks") or 0))
    pdf.cell(11, 4.8, f"{marks}M" if marks else "-", align="R")

    text = " ".join(step["representative_text"].split())
    text = text if len(text) <= 170 else text[:167].rsplit(" ", 1)[0] + "..."
    pdf.set_xy(text_x, y + 4.8)
    pdf.set_text_color(*INK_2)
    pdf.set_font("Helvetica", "", 7.8)
    pdf.multi_cell(text_w, 3.6, _safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    pdf.set_draw_color(*LINE)
    pdf.set_line_width(0.2)
    pdf.line(MARGIN, pdf.get_y() + 1.2, PAGE_W - MARGIN, pdf.get_y() + 1.2)
    pdf.ln(2.4)


def _draw_footer(pdf: FPDF, subject_code: str, site_url: str):
    pdf.set_auto_page_break(False)
    pdf.set_y(-13)
    pdf.set_draw_color(*LINE)
    pdf.set_line_width(0.2)
    pdf.line(MARGIN, pdf.get_y(), PAGE_W - MARGIN, pdf.get_y())
    pdf.ln(1.5)
    pdf.set_text_color(*INK_2)
    pdf.set_font("Helvetica", "", 7)
    left = f"Every topic and wording: {site_url.rstrip('/')}/s/{subject_code}" if site_url else "Papers Please"
    pdf.cell(CONTENT_W / 2, 4, _safe(left))
    pdf.cell(CONTENT_W / 2, 4, "Based on how often topics repeated - not a guarantee of what will appear.", align="R")


# ── Question bank (CSV and Excel) ─────────────────────────────────────────────
#
# Both are built from the snapshot JSON the website renders, so they always
# match the subject page: one row per printed question, grouped by topic.

BANK_COLUMNS = [
    # (header, Excel width, what it means: shown on the About sheet)
    ("Module", 9, "Module of the syllabus (1 to 5)."),
    ("Topic rank", 8, "1 = the topic most worth studying in that module."),
    ("Topic", 38, "Short name for the topic. Rows with the same topic asked the same thing."),
    ("How likely", 13, "Very likely: in 80% or more of the papers. Likely: in half or more. Occasional: fewer."),
    ("Asked in", 14, "How many of the analysed papers asked this topic."),
    ("Usual marks", 9, "Average marks the topic carried when it was asked."),
    ("Session", 13, "The exam the question comes from."),
    ("Question number", 11, "Where it was printed in that paper, e.g. Q3a."),
    ("Marks", 8, "Marks printed next to the question (blank if the paper did not print them)."),
    ("Question", 90, "The question as it was printed."),
]
BANK_HEADERS = [c[0] for c in BANK_COLUMNS]

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text) -> str:
    """One line of text, without control characters a spreadsheet would reject."""
    return " ".join(_CONTROL_CHARS.sub(" ", str(text or "")).split())


def likelihood(frequency_pct: float, total_papers: int) -> str:
    """The same label the subject page shows for a topic."""
    if total_papers <= 1:
        return "Only 1 paper"
    if frequency_pct >= 0.8:
        return "Very likely"
    if frequency_pct >= 0.5:
        return "Likely"
    return "Occasional"


def session_labels(data: dict) -> dict:
    """Session id → label, numbering repeats ("Model paper 1", "Model paper 2") so rows can be told apart."""
    by_label: dict[str, list[int]] = {}
    for s in data["sessions"]:
        by_label.setdefault(s["label"], []).append(s["id"])
    labels = {}
    for text, ids in by_label.items():
        for n, sid in enumerate(sorted(ids), start=1):   # lowest id = uploaded first
            labels[sid] = f"{text} {n}" if len(ids) > 1 else text
    return labels


def question_rows(data: dict) -> list[dict]:
    """One row per printed question: by module, topic rank, then newest session first."""
    total = data["paperCount"]
    order = {s["id"]: i for i, s in enumerate(data["sessions"])}   # newest first, like the site
    label = session_labels(data)
    rows = []
    for module in sorted(data["modules"], key=lambda m: m["no"]):
        for topic in sorted(module["topics"], key=lambda t: t["rank"]):
            common = {
                "Module": module["no"],
                "Topic rank": topic["rank"],
                "Topic": _clean(topic["label"]),
                "How likely": likelihood(topic["frequencyPct"], total),
                "Asked in": f"{topic['frequency']} of {total} paper{'s' if total != 1 else ''}",
                "Usual marks": int(round(topic["avgMarks"])) if topic.get("avgMarks") else None,
            }
            wordings = sorted(topic["wordings"], key=lambda w: (order.get(w["session"], len(order)), w["where"]))
            for w in wordings:
                rows.append({
                    **common,
                    "Session": label.get(w["session"], ""),
                    "Question number": w["where"],
                    "Marks": w["marks"],
                    "Question": _clean(w["text"]),
                })
    return rows


def _csv_cell(value):
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value   # never let a spreadsheet run a question as a formula
    return value


def generate_csv(data: dict) -> str:
    """
    The question bank as CSV, ready for Excel and Google Sheets on Windows:
    a UTF-8 byte-order mark (so Excel reads the text correctly), CRLF line
    endings, one header row and no title row, so sorting and filters just work.
    """
    buf = io.StringIO()
    buf.write("﻿")
    writer = csv.writer(buf, lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(BANK_HEADERS)
    for row in question_rows(data):
        writer.writerow([_csv_cell(row[h]) for h in BANK_HEADERS])
    return buf.getvalue()


# Excel colours (hex, no #): the same palette as the website and the cheat sheet
_XL_INK, _XL_INK_2, _XL_LINE, _XL_BAND = "1B2233", "4A5163", "DDD7C9", "F4F1EA"
_XL_LIKELY = {   # label → (fill, font), the site's hi / mid / lo pills
    "Very likely": ("F6E3E0", "8E2A20"),
    "Likely": ("F8E9D4", "7A4210"),
    "Occasional": ("ECE8DE", "3F4556"),
    "Only 1 paper": ("ECE8DE", "3F4556"),
}
_XL_CENTRED = {"Module", "Topic rank", "Usual marks", "Question number", "Marks"}


def _day(iso) -> str:
    """"2026-10-14T10:00:00+00:00" → "14 Oct 2026" (the same format as "Generated")."""
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d %b %Y")
    except ValueError:
        return ""


def generate_xlsx(data: dict) -> bytes:
    """
    The question bank as a formatted Excel workbook: an "All questions" sheet,
    one sheet per module and a short "About" sheet explaining the columns.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rows = question_rows(data)
    wb = Workbook()
    wb.properties.title = f"{data['code']} {data['name']} - question bank"
    wb.properties.creator = "Papers Please"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=_XL_INK)
    band_fill = PatternFill("solid", fgColor=_XL_BAND)
    rule = Border(bottom=Side(style="thin", color=_XL_LINE))
    top_left = Alignment(vertical="top")
    wrap = Alignment(vertical="top", wrap_text=True)
    centre = Alignment(vertical="top", horizontal="center")

    def question_sheet(ws, sheet_rows: list[dict], columns: list[tuple]):
        headers = [c[0] for c in columns]
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        banded, last_topic = True, None
        for r, row in enumerate(sheet_rows, start=2):
            topic = (row["Module"], row["Topic rank"])
            if topic != last_topic:          # alternate the fill per topic
                banded, last_topic = not banded, topic
            for c, header in enumerate(headers, start=1):
                value = row[header]
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"     # text stays text, even when it starts with "="
                cell.border = rule
                if header in ("Question", "Topic"):
                    cell.alignment = wrap
                else:
                    cell.alignment = centre if header in _XL_CENTRED else top_left
                if banded:
                    cell.fill = band_fill
                if header == "How likely" and value in _XL_LIKELY:
                    fill, font = _XL_LIKELY[value]
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(bold=True, color=font)
                elif header == "Topic":
                    cell.font = Font(bold=True, color=_XL_INK)

        for c, (_, width, _) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(sheet_rows) + 1, 2)}"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.active
    ws.title = "All questions"
    question_sheet(ws, rows, BANK_COLUMNS)
    without_module = [c for c in BANK_COLUMNS if c[0] != "Module"]
    for module_no in sorted({r["Module"] for r in rows}):
        question_sheet(wb.create_sheet(f"Module {module_no}"),
                       [r for r in rows if r["Module"] == module_no], without_module)

    # About: what the workbook covers and what each column means
    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 20, 90

    def about_row(key, value, bold_key=True):
        ws.append([key, value])
        if bold_key:
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    ws.append([f"{data['code']}  {_clean(data['name'])}"])
    ws["A1"].font = Font(bold=True, size=16, color=_XL_INK)
    ws.append(["Question bank from past VTU papers, grouped by topic."])
    ws["A2"].font = Font(italic=True, color=_XL_INK_2)
    ws.append([])
    about_row("Subject", f"{data['code']} - {_clean(data['name'])}")
    about_row("Papers analysed", data["paperCount"])
    labels = session_labels(data)
    about_row("Sessions covered", ", ".join(labels[s["id"]] for s in data["sessions"]))
    about_row("Questions", len(rows))
    about_row("Data updated", _day(data.get("updatedAt")))
    about_row("Generated", date.today().strftime("%d %b %Y"))
    ws.append([])
    ws.append(["Column", "What it means"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = header_font, header_fill
    for header, _, meaning in BANK_COLUMNS:
        about_row(header, meaning)
    ws.append([])
    about_row("Note", "Based on how often topics repeated in past papers - not a guarantee of what will appear. "
                      "Use the filter arrows on each sheet to pick a module, topic or session.")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
