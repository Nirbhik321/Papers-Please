"""Question bank downloads: a CSV that opens cleanly in Excel and a formatted .xlsx."""

import csv
import io
import json

import openpyxl
from sqlalchemy import delete, insert

from modules import exporter
from server.db import engine, subject_snapshots

CODE = "BCS998"
SNAPSHOT = {
    "code": CODE,
    "name": "Test Subject",
    "updatedAt": "2026-10-14T10:00:00+00:00",
    "paperCount": 3,
    "topicCount": 3,
    "moduleMarks": 20,
    "sessions": [
        {"id": 3, "type": "see", "label": "Jul 2025", "short": "Jul 25"},
        {"id": 2, "type": "see", "label": "Jan 2025", "short": "Jan 25"},
        {"id": 1, "type": "model", "label": "Model paper", "short": "Model"},
    ],
    "modules": [
        {"no": 2, "topics": [
            {"key": "c", "rank": 1, "label": "Formula topic", "text": "=1+1", "frequency": 1,
             "frequencyPct": 0.3333, "avgMarks": 6.0, "expectedMarks": 2.0, "sessions": [2],
             "wordings": [{"session": 2, "where": "Q3a", "marks": 6, "text": "=HYPERLINK(\"http://x\",\"click\")"}]},
        ]},
        {"no": 1, "topics": [
            {"key": "b", "rank": 2, "label": "Routing", "text": "Explain routing", "frequency": 2,
             "frequencyPct": 0.6667, "avgMarks": 8.0, "expectedMarks": 5.3, "sessions": [3, 1],
             "wordings": [
                 {"session": 1, "where": "Q2b", "marks": None, "text": "Explain routing, with an example"},
                 {"session": 3, "where": "Q1b", "marks": 8, "text": "Explain\n routing \"briefly\""},
             ]},
            {"key": "a", "rank": 1, "label": "TCP handshake", "text": "Explain the handshake", "frequency": 3,
             "frequencyPct": 1.0, "avgMarks": 10.0, "expectedMarks": 10.0, "sessions": [3, 2, 1],
             "wordings": [
                 {"session": 2, "where": "Q1a", "marks": 10, "text": "Explain the TCP 3-way handshake"},
                 {"session": 3, "where": "Q1a", "marks": 10, "text": "Describe the TCP handshake"},
                 {"session": 1, "where": "Q2a", "marks": 10, "text": "What is a handshake?"},
             ]},
        ]},
    ],
}


def _rows(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text.lstrip("﻿"), newline="")))


def test_csv_format():
    text = exporter.generate_csv(SNAPSHOT)
    assert text.startswith("﻿")
    assert "\r\n" in text and "\n" not in text.replace("\r\n", "")
    rows = _rows(text)
    assert rows[0] == exporter.BANK_HEADERS
    assert len(rows) == 1 + 6   # one row per printed question

    # Module, then topic rank, then newest session first
    order = [(r[0], r[1], r[6]) for r in rows[1:]]
    assert order == [("1", "1", "Jul 2025"), ("1", "1", "Jan 2025"), ("1", "1", "Model paper"),
                     ("1", "2", "Jul 2025"), ("1", "2", "Model paper"), ("2", "1", "Jan 2025")]

    first = dict(zip(rows[0], rows[1]))
    assert first == {"Module": "1", "Topic rank": "1", "Topic": "TCP handshake", "How likely": "Very likely",
                     "Asked in": "3 of 3 papers", "Usual marks": "10", "Session": "Jul 2025",
                     "Question number": "Q1a", "Marks": "10", "Question": "Describe the TCP handshake"}
    routing = dict(zip(rows[0], rows[4]))
    assert routing["How likely"] == "Likely" and routing["Question"] == 'Explain routing "briefly"'
    assert dict(zip(rows[0], rows[5]))["Marks"] == ""   # not printed on the paper
    assert dict(zip(rows[0], rows[6]))["How likely"] == "Occasional"


def test_csv_escapes_formulas():
    rows = _rows(exporter.generate_csv(SNAPSHOT))
    question = dict(zip(rows[0], rows[-1]))["Question"]
    assert question == "'=HYPERLINK(\"http://x\",\"click\")"


def test_xlsx_workbook():
    wb = openpyxl.load_workbook(io.BytesIO(exporter.generate_xlsx(SNAPSHOT)))
    assert wb.sheetnames == ["All questions", "Module 1", "Module 2", "About"]

    ws = wb["All questions"]
    assert [c.value for c in ws[1]] == exporter.BANK_HEADERS
    assert ws.max_row == 7
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:J7"
    assert ws["A1"].font.bold
    assert ws["J2"].alignment.wrap_text
    assert ws["D2"].value == "Very likely" and ws["D2"].fill.fgColor.rgb.endswith("F6E3E0")
    # The formula-like question is stored as plain text, not a formula
    assert ws["J7"].data_type == "s" and ws["J7"].value.startswith("=HYPERLINK")
    # Rows of the same topic share a fill; the next topic switches
    assert ws["A2"].fill.fgColor.rgb == ws["A4"].fill.fgColor.rgb != ws["A5"].fill.fgColor.rgb

    module_1 = wb["Module 1"]
    assert "Module" not in [c.value for c in module_1[1]]
    assert module_1.max_row == 6

    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True) if r[0]}
    assert about["Papers analysed"] == 3
    assert about["Sessions covered"] == "Jul 2025, Jan 2025, Model paper"
    assert about["Data updated"] == "14 Oct 2026"


def test_repeated_session_labels_are_numbered():
    # Two model-paper sets both read "Model paper": number them so their rows can be told apart
    data = {**SNAPSHOT, "sessions": SNAPSHOT["sessions"][:2] + [
        {"id": 1, "type": "model", "label": "Model paper", "short": "Model"},
        {"id": 4, "type": "model", "label": "Model paper", "short": "Model"},
    ]}
    assert exporter.session_labels(data) == {3: "Jul 2025", 2: "Jan 2025", 1: "Model paper 1", 4: "Model paper 2"}
    sessions = {r["Session"] for r in exporter.question_rows(data)}
    assert "Model paper 1" in sessions and "Model paper" not in sessions


def test_download_endpoints(client):
    with engine.begin() as conn:
        conn.execute(delete(subject_snapshots).where(subject_snapshots.c.subject_code == CODE))
        conn.execute(insert(subject_snapshots).values(
            subject_code=CODE, subject_name="Test Subject", paper_count=3, topic_count=3,
            data=json.dumps(SNAPSHOT), questions_csv="old format",
        ))
    try:
        r = client.get(f"/api/subjects/{CODE.lower()}/questions.csv")
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
        assert r.content.startswith(b"\xef\xbb\xbf")
        assert f'{CODE}_question_bank.csv' in r.headers["content-disposition"]
        assert _rows(r.content.decode("utf-8"))[0] == exporter.BANK_HEADERS

        r = client.get(f"/api/subjects/{CODE}/questions.xlsx")
        assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
        assert f'{CODE}_question_bank.xlsx' in r.headers["content-disposition"]
        assert "About" in openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames

        assert client.get("/api/subjects/BCS997/questions.xlsx").status_code == 404
    finally:
        with engine.begin() as conn:
            conn.execute(delete(subject_snapshots).where(subject_snapshots.c.subject_code == CODE))
